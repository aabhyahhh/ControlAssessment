"""
Lightweight text extraction for uploaded documents — docx/pdf/txt/xlsx only,
no OCR for v1. Shared by step 2 (SOP + workpaper ingestion) and, if needed
later, evidence-date extraction (timeline sufficiency also reads evidence
filenames/snippets, not full extraction).

xlsx support is here because monthly workpapers are commonly spreadsheets;
the engine reconciles their text content against the RCM row.
"""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger("engines.text_extraction")


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        return _extract_docx(file_path)
    if suffix == ".pdf":
        return _extract_pdf(file_path)
    if suffix == ".txt":
        return file_path.read_text(encoding="utf-8", errors="replace")
    if suffix in (".xlsx", ".xlsm"):
        return _extract_xlsx(file_path)
    raise ValueError(f"Unsupported document type '{suffix}'. Use .docx, .pdf, .txt, or .xlsx.")


def _extract_docx(file_path: Path) -> str:
    import docx

    document = docx.Document(str(file_path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def _extract_pdf(file_path: Path) -> str:
    import pdfplumber

    parts = []
    with pdfplumber.open(str(file_path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
    return "\n".join(parts)


def _extract_xlsx(file_path: Path) -> str:
    """Every non-empty cell across every sheet, row by row. Workpapers are
    tabular, so a flat "Sheet | col | col" dump is enough for the reconciliation
    prompt to see owners, dates, sign-offs and amounts."""
    from openpyxl import load_workbook

    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            parts.append(f"--- Sheet: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(parts)
