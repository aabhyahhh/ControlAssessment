"""
Excel exports for the redesigned flow.

  - build_phase_workbook: one workbook per step's result (editable shape).
  - build_final_report:    the multi-sheet deliverable covering all four steps.
  - build_rcm_workbook:    the working RCM (normalized + overrides).

The step-4 workbook IS the deliverable — a gap-assessment summary (received vs
expected, where the gap lies, severity), not a test-of-effectiveness
workpaper. The reports are universe-preserving: every control in the RCM
appears exactly once on the summary sheet.
"""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("engines.export_engine")

_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BODY_ALIGN = Alignment(vertical="top", wrap_text=True)
_THIN = Side(style="thin", color="B4C6E7")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _write_sheet(ws, headers: list[str], widths: list[int], rows: list[list[Any]]) -> None:
    for col, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col)].width = width

    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = _BODY_ALIGN
            cell.border = _BORDER
    ws.freeze_panes = "A2"


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def _join(values: Any) -> str:
    return ", ".join(str(v) for v in (values or []))


def _pct(value: Any) -> str:
    return "" if value is None else f"{round(value * 100)}%"


# ═══════════════════════════════════════════════════════════════════════════
#  Per-step workbook
# ═══════════════════════════════════════════════════════════════════════════


def build_phase_workbook(phase: int, controls: list[dict[str, Any]], result: dict[str, Any], dest_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active

    if phase == 1:
        ws.title = "Step 1 - RCM Intake"
        rows = [
            [
                c["control_id"],
                f"{round(c.get('completeness_pct', 0) * 100)}%",
                _join(c.get("missing_fields")),
                c.get("control_description", ""),
            ]
            for c in (result.get("controls") or [])
        ] or [[c["control_id"], "", "", c.get("control_description") or ""] for c in controls]
        _write_sheet(ws, ["Control ID", "Completeness", "Blank Fields", "Control Description"], [16, 14, 34, 60], rows)

    elif phase == 2:
        ws.title = "Step 2 - Adequacy"
        recon = {r["control_id"]: r for r in (result.get("reconciliation") or [])}
        align = {a["control_id"]: a for a in (result.get("control_alignment") or [])}
        wp = {w["control_id"]: w for w in (result.get("workpaper_coverage") or [])}
        _write_sheet(
            ws,
            ["Control ID", "Described in Docs", "Reconciliation", "SOP Alignment", "Mismatches",
             "Design Verdict", "Weak Dimensions", "Workpaper Months Missing"],
            [16, 16, 14, 16, 48, 18, 26, 30],
            [
                [
                    d["control_id"],
                    "yes" if (recon.get(d["control_id"]) or {}).get("described_in_docs") else "no",
                    _pct((recon.get(d['control_id']) or {}).get('reconciliation_pct')),
                    (align.get(d["control_id"]) or {}).get("alignment", ""),
                    "; ".join(
                        f"{m['field']}: RCM='{m['rcm_value']}' vs SOP='{m['sop_value']}'"
                        for m in (align.get(d["control_id"]) or {}).get("mismatches", [])
                    ),
                    d.get("verdict", ""),
                    _join(d.get("weak_dimensions")),
                    _join((wp.get(d["control_id"]) or {}).get("months_missing")),
                ]
                for d in (result.get("deficiencies") or [])
            ],
        )
        if result.get("coverage_gaps"):
            _write_sheet(
                wb.create_sheet("SOP Gaps"),
                ["SOP Step", "Description", "Coverage"],
                [14, 90, 14],
                [[g["sop_step_id"], g["description"], g["coverage"]] for g in result["coverage_gaps"]],
            )

    elif phase == 3:
        ws.title = "Step 3 - Evidence"
        _write_sheet(
            ws,
            ["Control ID", "Score", "Severity", "Expected Documents", "Received (uploaded)",
             "Declared not Uploaded", "Missing", "Extra Declared"],
            [16, 8, 12, 48, 40, 34, 34, 30],
            [
                [
                    r["control_id"], r.get("score", 0), r.get("severity") or "",
                    _join(r.get("required")), _join(r.get("matched")),
                    _join(r.get("declared_not_uploaded")), _join(r.get("missing")),
                    _join(r.get("extra_declared")),
                ]
                for r in (result.get("per_control") or [])
            ],
        )

    else:
        ws.title = "Step 4 - Gap Assessment"
        _write_sheet(
            ws,
            ["Control ID", "Severity", "Expected", "Received", "Missing Documents", "RCM Field Gaps",
             "SOP Alignment", "Reconciliation", "Design Verdict", "Workpaper Months Missing", "Summary"],
            [16, 12, 34, 34, 40, 30, 16, 14, 18, 30, 70],
            [
                [
                    r["control_id"], (r.get("severity") or "none").upper(),
                    _join(r.get("expected_documents")), _join(r.get("received_documents")),
                    _join(r.get("missing_documents")), _join(r.get("rcm_field_gaps")),
                    r.get("sop_alignment") or "",
                    _pct(r.get('reconciliation_pct')),
                    r.get("design_verdict") or "", _join(r.get("workpaper_months_missing")),
                    r.get("summary", ""),
                ]
                for r in (result.get("rows") or [])
            ],
        )
        rollup = (result.get("stats") or {}).get("severity_rollup") or {}
        _write_sheet(
            wb.create_sheet("Severity Rollup"),
            ["Severity", "Control Count"],
            [16, 16],
            [[k.upper(), v] for k, v in rollup.items()],
        )

    name = {1: "RCM_Intake", 2: "Adequacy", 3: "Evidence", 4: "Gap_Assessment"}[phase]
    dest = dest_dir / f"Step{phase}_{name}_{_timestamp()}.xlsx"
    wb.save(dest)
    return dest


# ═══════════════════════════════════════════════════════════════════════════
#  Working RCM
# ═══════════════════════════════════════════════════════════════════════════


def build_rcm_workbook(controls: list[dict[str, Any]], dest_dir: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "RCM"
    fields = [
        ("control_id", "Control ID"), ("control_description", "Control Description"),
        ("risk_description", "Risk Description"), ("risk_level", "Risk Level"),
        ("control_type", "Control Type"), ("control_nature", "Control Nature"),
        ("control_frequency", "Control Frequency"), ("control_owner", "Control Owner"),
        ("process", "Process"),
    ]
    _write_sheet(
        ws,
        [label for _, label in fields],
        [16, 55, 45, 12, 16, 16, 18, 20, 20],
        [[c.get(key) or "" for key, _ in fields] for c in sorted(controls, key=lambda x: x["control_id"])],
    )
    dest = dest_dir / f"RCM_Working_Copy_{_timestamp()}.xlsx"
    wb.save(dest)
    return dest


# ═══════════════════════════════════════════════════════════════════════════
#  Final multi-step report
# ═══════════════════════════════════════════════════════════════════════════


def build_final_report(
    project: dict[str, Any],
    controls: list[dict[str, Any]],
    phase_results: dict[int, dict[str, Any]],
    dest_dir: Path,
) -> Path:
    wb = Workbook()

    p1 = phase_results.get(1) or {}
    p2 = phase_results.get(2) or {}
    p3 = phase_results.get(3) or {}
    p4 = phase_results.get(4) or {}

    completeness = {c["control_id"]: c for c in (p1.get("controls") or [])}
    recon = {r["control_id"]: r for r in (p2.get("reconciliation") or [])}
    align = {a["control_id"]: a for a in (p2.get("control_alignment") or [])}
    deficiency = {d["control_id"]: d for d in (p2.get("deficiencies") or [])}
    workpaper = {w["control_id"]: w for w in (p2.get("workpaper_coverage") or [])}
    evidence = {r["control_id"]: r for r in (p3.get("per_control") or [])}
    gap = {r["control_id"]: r for r in (p4.get("rows") or [])}

    # ── Sheet 1: control universe ────────────────────────────────────
    universe_rows = []
    for c in sorted(controls, key=lambda x: x["control_id"]):
        cid = c["control_id"]
        g = gap.get(cid, {})
        ev = evidence.get(cid, {})
        universe_rows.append([
            cid,
            c.get("control_description") or "",
            f"{round(completeness.get(cid, {}).get('completeness_pct', 0) * 100)}%",
            "yes" if (recon.get(cid) or {}).get("described_in_docs") else "no",
            (align.get(cid) or {}).get("alignment", ""),
            (deficiency.get(cid) or {}).get("verdict", ""),
            f"{ev.get('score', 0)}%",
            _join((workpaper.get(cid) or {}).get("months_missing")),
            (g.get("severity") or "none").upper(),
            g.get("summary", ""),
        ])

    ws = wb.active
    ws.title = "Control Universe"
    _write_sheet(
        ws,
        ["Control ID", "Control Description", "RCM Completeness", "Described in Docs", "SOP Alignment",
         "Design Verdict", "Evidence Score", "Workpaper Months Missing", "Gap Severity", "Summary"],
        [16, 46, 16, 16, 16, 18, 14, 28, 14, 64],
        universe_rows,
    )

    # ── Sheet 2: Step 1 ─────────────────────────────────────────────
    _write_sheet(
        wb.create_sheet("Step 1 - RCM Intake"),
        ["Control ID", "Completeness", "Blank Fields", "Control Description"],
        [16, 14, 34, 60],
        [
            [c["control_id"], f"{round(c.get('completeness_pct', 0) * 100)}%", _join(c.get("missing_fields")),
             c.get("control_description", "")]
            for c in (p1.get("controls") or [])
        ],
    )

    # ── Sheet 3: Step 2 ─────────────────────────────────────────────
    _write_sheet(
        wb.create_sheet("Step 2 - Adequacy"),
        ["Control ID", "Described in Docs", "Reconciliation", "SOP Alignment", "Mismatches",
         "Design Verdict", "Weak Dimensions", "Workpaper Months Missing"],
        [16, 16, 14, 16, 48, 18, 26, 30],
        [
            [
                d["control_id"],
                "yes" if (recon.get(d["control_id"]) or {}).get("described_in_docs") else "no",
                _pct((recon.get(d['control_id']) or {}).get('reconciliation_pct')),
                (align.get(d["control_id"]) or {}).get("alignment", ""),
                "; ".join(
                    f"{m['field']}: RCM='{m['rcm_value']}' vs SOP='{m['sop_value']}'"
                    for m in (align.get(d["control_id"]) or {}).get("mismatches", [])
                ),
                d.get("verdict", ""), _join(d.get("weak_dimensions")),
                _join((workpaper.get(d["control_id"]) or {}).get("months_missing")),
            ]
            for d in (p2.get("deficiencies") or [])
        ],
    )
    if p2.get("coverage_gaps"):
        _write_sheet(
            wb.create_sheet("Step 2 - SOP Gaps"),
            ["SOP Step", "Description", "Coverage"],
            [14, 90, 14],
            [[g["sop_step_id"], g["description"], g["coverage"]] for g in p2["coverage_gaps"]],
        )

    # ── Sheet 4: Step 3 ─────────────────────────────────────────────
    _write_sheet(
        wb.create_sheet("Step 3 - Evidence"),
        ["Control ID", "Score", "Severity", "Expected Documents", "Received", "Declared not Uploaded", "Missing"],
        [16, 8, 12, 48, 40, 34, 34],
        [
            [
                r["control_id"], r.get("score", 0), r.get("severity") or "",
                _join(r.get("required")), _join(r.get("matched")),
                _join(r.get("declared_not_uploaded")), _join(r.get("missing")),
            ]
            for r in (p3.get("per_control") or [])
        ],
    )

    # ── Sheet 5: Step 4 gap assessment ─────────────────────────────
    _write_sheet(
        wb.create_sheet("Step 4 - Gap Assessment"),
        ["Control ID", "Severity", "Expected", "Received", "Missing Documents", "RCM Field Gaps",
         "SOP Alignment", "Reconciliation", "Design Verdict", "Workpaper Months Missing", "Summary"],
        [16, 12, 34, 34, 40, 30, 16, 14, 18, 30, 70],
        [
            [
                r["control_id"], (r.get("severity") or "none").upper(),
                _join(r.get("expected_documents")), _join(r.get("received_documents")),
                _join(r.get("missing_documents")), _join(r.get("rcm_field_gaps")),
                r.get("sop_alignment") or "",
                _pct(r.get('reconciliation_pct')),
                r.get("design_verdict") or "", _join(r.get("workpaper_months_missing")),
                r.get("summary", ""),
            ]
            for r in (p4.get("rows") or [])
        ],
    )
    rollup = (p4.get("stats") or {}).get("severity_rollup") or {}
    if rollup:
        _write_sheet(
            wb.create_sheet("Step 4 - Severity Rollup"),
            ["Severity", "Control Count"],
            [16, 16],
            [[k.upper(), v] for k, v in rollup.items()],
        )

    safe_name = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in (project.get("name") or "project"))[:60]
    dest = dest_dir / f"Gap_Assessment_Report_{safe_name}_{_timestamp()}.xlsx"
    wb.save(dest)
    return dest
